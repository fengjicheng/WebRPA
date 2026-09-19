# -*- coding: utf-8 -*-
"""3.2.6 用户反馈回归：小数等待、OCR 回退、Excel 句柄/空尾与多层 iframe。"""
import sys
from types import SimpleNamespace

import pytest


async def test_fixed_wait_accepts_decimal_seconds(monkeypatch, make_context):
    from app.executors.basic import WaitExecutor

    slept = []

    async def fake_sleep(seconds):
        slept.append(seconds)

    monkeypatch.setattr("app.executors.basic.asyncio.sleep", fake_sleep)
    result = await WaitExecutor().execute({"waitType": "time", "duration": 0.25}, make_context())
    assert result.success
    assert slept == [0.25]


@pytest.mark.parametrize(
    ("value", "target", "expected"),
    [
        ("12", "integer", 12),
        ("12.9", "integer", 12),
        ("12.5", "float", 12.5),
        ("是", "boolean", True),
        ("false", "boolean", False),
        ("[1, 2]", "list", [1, 2]),
        ('{"a": 1}', "dictionary", {"a": 1}),
        ({"a": 1}, "string", '{"a": 1}'),
    ],
)
async def test_type_convert_all_supported_types(value, target, expected, make_context):
    from app.executors.basic_variable import TypeConvertExecutor

    ctx = make_context({"source": value})
    result = await TypeConvertExecutor().execute(
        {"inputValue": "{source}", "targetType": target, "resultVariable": "out"}, ctx
    )
    assert result.success, result.error
    assert ctx.get_variable("out") == expected


async def test_read_excel_trims_only_trailing_edited_empty_cells(tmp_path, make_context):
    import openpyxl
    from openpyxl.styles import PatternFill
    from app.executors.advanced import ReadExcelExecutor

    path = tmp_path / "edited_cells.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "header"
    ws["A2"] = "first"
    ws["A3"] = None  # 数据中间的空值必须保留
    ws["A4"] = "last"
    ws["A500"].fill = PatternFill("solid", fgColor="FFFF00")  # 曾编辑/仅格式化的尾部空格
    wb.save(path)
    wb.close()

    ctx = make_context()
    result = await ReadExcelExecutor().execute(
        {
            "fileName": str(path),
            "readMode": "column",
            "columnIndex": "A",
            "startRow": 2,
            "variableName": "rows",
        },
        ctx,
    )
    assert result.success, result.error
    assert ctx.get_variable("rows") == ["first", None, "last"]


async def test_excel_create_closes_handle_and_can_run_twice(tmp_path, make_context):
    import openpyxl
    from app.executors.advanced_openpyxl import ExcelCreateExecutor

    path = tmp_path / "repeat.xlsx"
    executor = ExcelCreateExecutor()
    config = {
        "filePath": str(path),
        "sheetNames": "数据,汇总",
        "overwrite": True,
        "openAfterCreate": False,
    }
    first = await executor.execute(config, make_context())
    second = await executor.execute(config, make_context())
    assert first.success, first.error
    assert second.success, second.error
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        assert wb.sheetnames == ["数据", "汇总"]
    finally:
        wb.close()
    assert [p.name for p in tmp_path.iterdir()] == ["repeat.xlsx"]


async def test_excel_close_executor_closes_selected_workbook(monkeypatch, tmp_path, make_context):
    from app.executors import advanced_openpyxl as module

    path = tmp_path / "opened.xlsx"
    calls = []

    def fake_close(file_path, close_all, save_changes):
        calls.append((file_path, close_all, save_changes))
        return 1

    monkeypatch.setattr(module, "_close_excel_workbooks", fake_close)
    result = await module.ExcelCloseExecutor().execute(
        {"filePath": str(path), "closeAll": False, "saveChanges": True}, make_context()
    )
    assert result.success, result.error
    assert result.data["closed"] == 1
    assert calls == [(str(path), False, True)]


def test_general_ocr_falls_back_when_easyocr_is_unavailable(monkeypatch):
    from app.executors import media_recognition as module

    monkeypatch.setattr(module, "get_easyocr_reader", lambda: (_ for _ in ()).throw(ImportError("missing")))

    class FakeRapidOCR:
        def __call__(self, _image):
            return [
                ([[0, 20], [1, 20], [1, 21], [0, 21]], "第二行", 0.9),
                ([[0, 1], [1, 1], [1, 2], [0, 2]], "第一行", 0.9),
            ], 0.01

    monkeypatch.setattr(module, "_rapidocr_reader", FakeRapidOCR())
    text, engine = module._read_general_ocr(object())
    assert engine == "RapidOCR"
    assert text == "第一行\n第二行"


async def test_picker_injects_every_existing_nested_frame(monkeypatch):
    from app.services import browser_manager
    from app.services import browser_engine

    class FakeFrame:
        def __init__(self, url):
            self.url = url
            self.scripts = []

        async def evaluate(self, script):
            self.scripts.append(script)
            return None

    class FakePage(FakeFrame):
        def __init__(self, frames):
            super().__init__("https://top.test")
            self.frames = frames
            self.main_frame = frames[0]

    class FakeContext:
        def __init__(self, pages):
            self.pages = pages
            self.init_scripts = []

        async def add_init_script(self, script):
            self.init_scripts.append(script)

    frames = [FakeFrame("https://top.test"), FakeFrame("https://level1.test"), FakeFrame("https://level2.test")]
    ctx = FakeContext([FakePage(frames)])
    monkeypatch.setattr(browser_engine, "get_context", lambda: ctx)
    monkeypatch.setattr(browser_manager, "_picker_init_script_registered", False)
    monkeypatch.setattr(browser_manager, "_picker_init_ctx_id", None)

    result = await browser_manager._start_picker_engine()
    assert result["success"]
    assert all(browser_manager.PICKER_SCRIPT in frame.scripts for frame in frames)


async def test_switch_iframe_can_descend_two_levels(make_context):
    from app.executors.basic import SwitchIframeExecutor

    class FakeFrame:
        def __init__(self, name, url, children=None):
            self.name = name
            self.url = url
            self.child_frames = children or []

        def is_detached(self):
            return False

        async def wait_for_load_state(self, *_args, **_kwargs):
            return None

    level2 = FakeFrame("level2", "https://level2.test")
    level1 = FakeFrame("level1", "https://level1.test", [level2])
    root_frame = FakeFrame("", "https://top.test", [level1])

    class FakePage:
        url = "https://top.test"
        main_frame = root_frame

    ctx = make_context()
    ctx.page = FakePage()
    executor = SwitchIframeExecutor()

    first = await executor.execute({"locateBy": "index", "iframeIndex": 0}, ctx)
    second = await executor.execute({"locateBy": "index", "iframeIndex": 0}, ctx)

    assert first.success, first.error
    assert second.success, second.error
    assert ctx.page is level2
    assert ctx._iframe_locator_path == [
        {"type": "index", "value": 0},
        {"type": "index", "value": 0},
    ]
